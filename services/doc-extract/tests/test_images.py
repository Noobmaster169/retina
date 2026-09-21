import io

from conftest import fixture_bytes
from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image, ImageDraw

from app import create_app, extract_bytes
from registry import format_of
from storage import MemoryStorage

"""
The case this all exists for: a document whose content is a picture.

Before, a bill of lading pasted into Word read as its covering sentence and
nothing else, at `unreadable: false`, so the pipeline typed it, extracted no
fields from it and reported a missing value. Nothing anywhere said a picture had
been dropped.
"""


def picture(width: int = 1000, height: int = 600, fmt: str = "PNG") -> bytes:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    for line, text in enumerate(["BILL OF LADING", "CONSIGNEE: MOORIM SP CO LTD", "POD: CALLAO, PERU"]):
        draw.text((40, 40 + line * 90), text, fill=(0, 0, 0))
    buffer = io.BytesIO()
    image.save(buffer, fmt)
    return buffer.getvalue()


def docx_with(text: str | None, image: bytes | None) -> bytes:
    document = Document()
    if text:
        document.add_paragraph(text)
    if image:
        document.add_picture(io.BytesIO(image), width=Inches(6))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_with(rows: list[tuple[str, str]], image: bytes | None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for label, value in rows:
        sheet.append([label, value])
    if image:
        sheet.add_image(SheetImage(io.BytesIO(image)), "A5")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_a_covering_sentence_over_a_pasted_document_is_not_a_readable_document():
    """The silent case. Its text clears the floor, so only the picture says otherwise."""
    data = docx_with("Dear Sir, please find attached the draft bill of lading.", picture())
    found = extract_bytes(data, "bl.docx", None)

    assert found.response.has_images is True
    assert len(found.extracted.images) == 1
    assert found.extracted.images[0].origin == "embedded"
    assert any("picture" in warning for warning in found.response.warnings)


def test_a_word_document_with_no_picture_is_read_as_before():
    found = extract_bytes(fixture_bytes("email_055_BL.docx"), "email_055_BL.docx", None)
    assert found.response.has_images is False
    assert found.extracted.images == []
    assert found.response.unreadable is False


def test_a_spreadsheet_carries_its_pictures_too():
    data = xlsx_with([("Shipping Instruction", ""), ("Reference", "5RSG-00133")], picture())
    found = extract_bytes(data, "si.xlsx", None)

    assert found.response.has_images is True
    assert found.extracted.images[0].origin == "embedded"
    # The cells it could read are still read; the picture is what it could not.
    assert "5RSG-00133" in found.response.text


def test_furniture_is_not_worth_a_look():
    """A logo or a signature strip is not a document, and reading each would cost a call."""
    found = extract_bytes(docx_with("Regards, Willy", picture(width=40, height=20)), "note.docx", None)
    assert found.response.has_images is False


def test_a_photographed_document_is_a_document():
    """An image attachment used to be an unknown format and was never looked at."""
    found = extract_bytes(picture(), "bl_photo.jpg", None)

    assert found.response.format == "image"
    assert found.response.unreadable is False, "a legible photograph is not unreadable"
    assert len(found.extracted.images) == 1
    assert found.extracted.images[0].origin == "page"


def test_a_multi_page_fax_gives_a_page_each():
    pages = [Image.new("RGB", (800, 600), "white") for _ in range(3)]
    buffer = io.BytesIO()
    pages[0].save(buffer, "TIFF", save_all=True, append_images=pages[1:])

    found = extract_bytes(buffer.getvalue(), "fax.tif", None)

    assert found.response.format == "image"
    assert [image.index for image in found.extracted.images] == [1, 2, 3]


def test_the_bytes_decide_the_format_not_the_name():
    """Real mail carries a PDF named .docx and a TIFF named .pdf."""
    assert format_of("bl.docx", None, fixture_bytes("email_059_SI.pdf")) == "pdf"
    assert format_of("scan.pdf", None, picture(fmt="TIFF")) == "image"
    # With nothing to go on, the name still decides.
    assert format_of("notes.txt", None, b"") == "txt"


def test_the_pixels_are_written_where_the_caller_asked():
    storage = MemoryStorage()
    client = TestClient(create_app(storage))
    storage.put("a/bl.docx", docx_with("Dear Sir,", picture()), None)

    body = client.post(
        "/extract",
        json={"key": "a/bl.docx", "filename": "bl.docx", "content_type": None, "out_prefix": "a/images/bl.docx"},
    ).json()

    assert [image["key"] for image in body["images"]] == ["a/images/bl.docx/1.png"]
    assert storage.objects["a/images/bl.docx/1.png"][1] == "image/png"
    assert body["images"][0]["origin"] == "embedded"


def test_without_somewhere_to_put_them_the_caller_is_still_told():
    """A caller that asked for no prefix must not be left thinking it saw the whole document."""
    storage = MemoryStorage()
    client = TestClient(create_app(storage))
    storage.put("a/bl.docx", docx_with("Dear Sir,", picture()), None)

    body = client.post(
        "/extract", json={"key": "a/bl.docx", "filename": "bl.docx", "content_type": None}
    ).json()

    assert body["has_images"] is True and body["images"] == []
    assert any("not read" in warning for warning in body["warnings"])
