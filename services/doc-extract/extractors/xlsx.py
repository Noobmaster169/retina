from io import BytesIO

from openpyxl import load_workbook

from extractors.base import Extracted, ExtractedPage, cell_text
from extractors.media import images_in_package


def _row_text(values: list[str]) -> str:
    """`A: B` when both are set, `A` alone when only the first is; further cells follow with ` | `."""
    if not any(values):
        return ""
    first, rest = values[0], [v for v in values[1:] if v]
    if not rest:
        return first
    joined = " | ".join(rest)
    return f"{first}: {joined}" if first else joined


def extract_xlsx(data: bytes) -> Extracted:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    pages: list[ExtractedPage] = []
    try:
        # One page per sheet: a workbook with several is several documents' worth of text.
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines = []
            for row in sheet.iter_rows(values_only=True):
                text = _row_text([cell_text(value) for value in row])
                if text:
                    lines.append(text)
            pages.append(ExtractedPage(index=index, text="\n".join(lines), source="text_layer"))
    finally:
        workbook.close()
    images, warnings = images_in_package(data, "xl", start=len(pages) + 1)
    return Extracted(pages=pages, images=images, warnings=warnings)
